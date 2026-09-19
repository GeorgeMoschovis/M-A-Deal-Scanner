"""The Word export follows the house format: Calibri 10, justified, bold/italic rules, links."""

import io
import zipfile

import docx
import pytest
from docx.enum.text import WD_ALIGN_PARAGRAPH

import docx_export
from helpers import make_deal


@pytest.fixture
def document_bytes():
    full = make_deal(
        "Kyklos Participations", "Akritas", "Kyklos Participations Reaches 95.96% Stake in Akritas",
        "https://www.euro2day.gr/a", industry_header="wood processing & building materials",
        deal_financials=[{"label": "Implied EV", "value": "€45.2m"}],
        deal_footnotes=["Implied equity value plus FY-24 net debt"],
        target={"description": "Akritas makes wood panels.",
                "financials": [{"label": "Revenue – 2024", "value": "€38.3m"}]},
        buyer={"description": "Kyklos is a holding company.", "financials": []})
    full["sources"].append({"title": "t", "url": "https://www.capital.gr/b", "name": "n"})
    thin = make_deal("PPC", "ABO Energy", "PPC Completes ABO Energy Acquisition",
                     "https://www.businessdaily.gr/c", industry_header="energy",
                     financials_note="financials not disclosed")
    rumour = make_deal("", "Westinghouse", "Westinghouse Explores IPO", "https://www.powergame.gr/d",
                       item_type="rumour", industry_header="tech")
    return docx_export.build_docx([full, thin, rumour])


@pytest.fixture
def document(document_bytes):
    return docx.Document(io.BytesIO(document_bytes))


def test_normal_style_is_calibri_10_justified(document):
    normal = document.styles["Normal"]
    assert normal.font.name == "Calibri"
    assert normal.font.size.pt == 10
    assert normal.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY


def test_every_paragraph_and_run_is_justified_calibri_10(document):
    for paragraph in document.paragraphs:
        assert paragraph.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
        for run in paragraph.runs:
            assert run.font.name == "Calibri"
            assert run.font.size.pt == 10


def find(document, text):
    return next(p for p in document.paragraphs if p.text.startswith(text))


def test_headlines_and_industry_headers_are_bold_and_headers_are_upper_case(document):
    for text in ("Kyklos Participations Reaches", "PPC Completes", "WOOD PROCESSING", "ENERGY",
                 "The Target", "The Buyer", "RUMOURS AND DEVELOPMENTS"):
        assert all(run.bold for run in find(document, text).runs), text
    headers = [p.text for p in document.paragraphs if p.text.isupper()]
    assert headers == ["ENERGY", "WOOD PROCESSING & BUILDING MATERIALS", "RUMOURS AND DEVELOPMENTS"]


def test_footnotes_and_the_undisclosed_note_are_italic(document):
    for text in ("*Implied equity value", "financials not disclosed"):
        assert all(run.italic for run in find(document, text).runs), text


def test_body_text_is_not_bold_or_italic(document):
    body = find(document, "Kyklos Participations did something")
    assert not any(run.bold or run.italic for run in body.runs)


def test_financial_lines_use_label_colon_value(document):
    assert find(document, "Implied EV").text == "Implied EV: €45.2m"
    assert find(document, "Revenue – 2024").text == "Revenue – 2024: €38.3m"


def test_every_entry_ends_with_clickable_source_links(document, document_bytes):
    assert find(document, "Source: euro2day.gr, capital.gr")
    xml = zipfile.ZipFile(io.BytesIO(document_bytes)).read("word/document.xml").decode("utf8")
    assert xml.count("<w:hyperlink ") == 4
    rels = zipfile.ZipFile(io.BytesIO(document_bytes)).read("word/_rels/document.xml.rels").decode("utf8")
    for url in ("https://www.euro2day.gr/a", "https://www.capital.gr/b",
                "https://www.businessdaily.gr/c", "https://www.powergame.gr/d"):
        assert url in rels


def test_rumours_come_last_and_carry_no_financial_block(document):
    texts = [p.text for p in document.paragraphs]
    start = texts.index("RUMOURS AND DEVELOPMENTS")
    assert texts[start + 1] == "Westinghouse Explores IPO"
    assert not any(":" in t and "€" in t for t in texts[start:])


def test_empty_digest_still_produces_a_readable_file():
    empty = docx.Document(io.BytesIO(docx_export.build_docx([])))
    assert [p.text for p in empty.paragraphs] == ["No matching deals found."]
