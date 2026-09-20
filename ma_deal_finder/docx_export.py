"""Word export: Calibri 10, justified, bold headlines, bold all-caps industry headers, italic footnotes."""

from datetime import datetime, timezone
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from .extract import group_deals, source_links

FONT = "Calibri"
SIZE_PT = 10


def _set_fonts(rfonts):
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), FONT)


def _new_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = section.right_margin = Cm(2.2)
    section.top_margin = section.bottom_margin = Cm(2.2)
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(SIZE_PT)
    _set_fonts(normal.element.get_or_add_rPr().get_or_add_rFonts())
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    props = doc.core_properties
    props.title = "M&A digest"
    props.author = props.last_modified_by = "M&A Deal Finder"
    props.comments = ""
    props.created = props.modified = datetime.now(timezone.utc).replace(tzinfo=None)
    return doc


def _run(paragraph, text, bold=False, italic=False):
    run = paragraph.add_run(text)
    run.bold, run.italic = bold, italic
    run.font.name = FONT
    run.font.size = Pt(SIZE_PT)
    _set_fonts(run._r.get_or_add_rPr().get_or_add_rFonts())


def _paragraph(doc, runs, after=6, keep_next=False):
    """runs is a list of (text, bold, italic)."""
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt = paragraph.paragraph_format
    fmt.space_before, fmt.space_after, fmt.keep_with_next = Pt(0), Pt(after), keep_next
    for text, bold, italic in runs:
        _run(paragraph, text, bold, italic)
    return paragraph


def _add_hyperlink(paragraph, url, text):
    rel_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    _set_fonts(fonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), str(SIZE_PT * 2))
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.extend([fonts, color, size, underline])
    label = OxmlElement("w:t")
    label.text = text
    label.set(qn("xml:space"), "preserve")
    run.extend([props, label])
    link.append(run)
    paragraph._p.append(link)


def _lines_block(doc, lines):
    for i, line in enumerate(lines):
        _paragraph(doc, [(line, False, False)], after=6 if i == len(lines) - 1 else 0)


def _party(doc, title, party):
    _paragraph(doc, [(title, True, False)], after=2, keep_next=True)
    _paragraph(doc, [(party["description"], False, False)], after=3 if party["financials"] else 6)
    _lines_block(doc, [f'{f["label"]}: {f["value"]}' for f in party["financials"]])


def _sources(doc, deal):
    paragraph = _paragraph(doc, [("Source: ", False, False)], after=12)
    links = source_links(deal)
    for i, (label, url) in enumerate(links):
        if url.startswith(("http://", "https://")):
            _add_hyperlink(paragraph, url, label)
        else:
            _run(paragraph, label)
        if i < len(links) - 1:
            _run(paragraph, ", ")


def _deal(doc, deal):
    _paragraph(doc, [(deal["headline"], True, False)], after=3, keep_next=True)
    _paragraph(doc, [(deal["description"], False, False)])
    _lines_block(doc, [f'{f["label"]}: {f["value"]}' for f in deal["deal_financials"]])
    for note in deal["deal_footnotes"]:
        _paragraph(doc, [(note, False, True)], after=3)
    if deal["financials_note"]:
        _paragraph(doc, [(deal["financials_note"], False, True)], after=6)
    if deal["target"]:
        _party(doc, "The Target", deal["target"])
    if deal["buyer"]:
        _party(doc, "The Buyer", deal["buyer"])
    _sources(doc, deal)


def _development(doc, item):
    _paragraph(doc, [(item["headline"], True, False)], after=3, keep_next=True)
    _paragraph(doc, [(item["description"], False, False)])
    _sources(doc, item)


def build_docx(deals):
    """Returns the digest as .docx bytes."""
    doc = _new_document()
    groups, developments = group_deals(deals)
    if not groups and not developments:
        _paragraph(doc, [("No matching deals found.", False, False)])
    for header, items in groups:
        _paragraph(doc, [(header.upper(), True, False)], after=6, keep_next=True)
        for deal in items:
            _deal(doc, deal)
    if developments:
        _paragraph(doc, [("RUMOURS AND DEVELOPMENTS", True, False)], after=6, keep_next=True)
        for item in developments:
            _development(doc, item)
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
